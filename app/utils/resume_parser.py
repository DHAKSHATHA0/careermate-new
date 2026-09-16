import os
import mimetypes
from werkzeug.utils import secure_filename
import pdfplumber
from docx import Document

class ResumeParser:
    """Parse resumes from PDF and DOCX files"""
    
    ALLOWED_EXTENSIONS = {'pdf', 'docx'}
    ALLOWED_MIMETYPES = {'application/pdf', 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'}
    MAX_FILE_SIZE = 16 * 1024 * 1024  # 16MB
    
    @staticmethod
    def validate_file(file):
        """Validate file before processing"""
        errors = []
        
        # Check if file exists
        if not file or file.filename == '':
            errors.append('No file selected')
            return False, errors
        
        # Check file extension
        filename = secure_filename(file.filename)
        ext = filename.rsplit('.', 1)[1].lower() if '.' in filename else ''
        
        if ext not in ResumeParser.ALLOWED_EXTENSIONS:
            errors.append(f'Invalid file type. Allowed: {", ".join(ResumeParser.ALLOWED_EXTENSIONS)}')
        
        # Check MIME type
        mime_type, _ = mimetypes.guess_type(filename)
        if mime_type not in ResumeParser.ALLOWED_MIMETYPES:
            errors.append('Invalid MIME type. Only PDF and DOCX allowed.')
        
        # Check file size
        file.seek(0, os.SEEK_END)
        file_size = file.tell()
        file.seek(0)
        
        if file_size > ResumeParser.MAX_FILE_SIZE:
            errors.append(f'File too large. Maximum size: 16MB')
        
        if file_size == 0:
            errors.append('File is empty')
        
        return len(errors) == 0, errors
    
    @staticmethod
    def parse_pdf(file_path):
        """Extract text from PDF"""
        text = ""
        try:
            with pdfplumber.open(file_path) as pdf:
                for page in pdf.pages:
                    text += page.extract_text() or ""
        except Exception as e:
            raise Exception(f"Error parsing PDF: {str(e)}")
        
        return text
    
    @staticmethod
    def parse_docx(file_path):
        """Extract text from DOCX"""
        text = ""
        try:
            doc = Document(file_path)
            for para in doc.paragraphs:
                text += para.text + "\n"
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        text += cell.text + " "
        except Exception as e:
            raise Exception(f"Error parsing DOCX: {str(e)}")
        
        return text
    
    @staticmethod
    def extract_text(file_path):
        """Extract text from resume file"""
        ext = file_path.rsplit('.', 1)[1].lower()
        
        if ext == 'pdf':
            return ResumeParser.parse_pdf(file_path)
        elif ext == 'docx':
            return ResumeParser.parse_docx(file_path)
        else:
            raise ValueError(f"Unsupported file type: {ext}")
